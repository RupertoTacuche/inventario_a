"""Exporta el reporte diario a un archivo .xlsx (openpyxl, Python puro: funciona en el APK)."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from model.entities import DailyReport

from .money import IVA_PERCENT

NAVY = "0A1A33"
MONEY = '"$"#,##0.00;[Red]-"$"#,##0.00'


def build_daily_xlsx(r: DailyReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen del día"
    bold = Font(bold=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    day = date.fromisoformat(r.day).strftime("%d/%m/%Y")
    ws["A1"] = f"RESUMEN DE VENTAS DEL DÍA ({day})"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    row = 3
    ws.cell(row, 1, "Métricas clave del día").font = bold
    for label, value, fmt in (("Total vendido (ingresos)", r.total_cents / 100, MONEY),
                              ("Total ganancia (neto)", r.net_profit_cents / 100, MONEY),
                              ("Número de transacciones", r.transactions, "0")):
        row += 1
        ws.cell(row, 1, label)
        cell = ws.cell(row, 2, value)
        cell.number_format = fmt

    row += 2
    ws.cell(row, 1, "Productos vendidos y ganancia").font = bold
    row += 1
    for col, title in enumerate(("Producto", "Cantidad", "Importe", "Ganancia"), 1):
        c = ws.cell(row, col, title)
        c.font, c.fill, c.border = head_font, head_fill, box
        c.alignment = Alignment(horizontal="center")
    for p in r.products:
        row += 1
        values = (p.name, p.quantity, p.revenue_cents / 100, p.profit_cents / 100)
        for col, v in enumerate(values, 1):
            c = ws.cell(row, col, v)
            c.border = box
            if col >= 3:
                c.number_format = MONEY
    if not r.products:
        row += 1
        ws.cell(row, 1, "Sin ventas registradas.")

    row += 2
    ws.cell(row, 1, "Resumen financiero diario").font = bold
    lines = (("Unidades vendidas", r.units_sold, "0"),
             ("Subtotal de ventas brutas", r.subtotal_cents / 100, MONEY),
             (f"IVA ({IVA_PERCENT}%) colectado", r.iva_cents / 100, MONEY),
             ("Descuentos totales", -r.discount_cents / 100, MONEY),
             ("INGRESOS TOTALES (ventas + IVA − desc.)", r.total_cents / 100, MONEY),
             ("Costo de lo vendido", r.cost_cents / 100, MONEY),
             ("Ganancia bruta (ventas − costo)", r.gross_profit_cents / 100, MONEY),
             ("TOTAL NETO GANADO", r.net_profit_cents / 100, MONEY))
    for label, value, fmt in lines:
        row += 1
        ws.cell(row, 1, label)
        c = ws.cell(row, 2, value)
        c.number_format = fmt
        if label.isupper() or label.startswith("INGRESOS") or label.startswith("TOTAL"):
            ws.cell(row, 1).font = bold
            c.font = bold

    ws.column_dimensions["A"].width = 44
    for col in "BCD":
        ws.column_dimensions[col].width = 16
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

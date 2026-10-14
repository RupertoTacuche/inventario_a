"""Pruebas de humo de la vista con una página simulada (no abre ventana)."""
import asyncio
import types

import pytest

ft = pytest.importorskip("flet")

from view import compat as c  # noqa: E402
from view.app_view import build_app  # noqa: E402


class FakeClipboard:
    def __init__(self):
        self.value = None

    async def set(self, text):
        self.value = text


class FakePage:
    def __init__(self, width=420, platform=None):
        self.width = width
        self.platform = platform or ft.PagePlatform.WINDOWS
        self.services, self.dialogs, self.added = [], [], []
        self.window = types.SimpleNamespace()
        self.clipboard = FakeClipboard()
        self.on_resize = None

    def update(self): pass
    def add(self, *controls): self.added.extend(controls)
    def show_dialog(self, d): self.dialogs.append(d)
    def pop_dialog(self): self.dialogs.pop()


def snack_text(page):
    last = page.dialogs[-1]
    return last.content.value if isinstance(last, ft.SnackBar) else None


@pytest.fixture
def ui(app):
    page = FakePage()
    build_app(page, app.inventory, app.sales_ctl, app.report)
    # recuperamos las páginas a través de los handlers de Tabs
    tabs = page.added[0].content.controls[0].content
    return page, tabs, app


def _pages(app, page):
    from view.pages.inventory_page import InventoryPage
    from view.pages.report_page import ReportPage
    from view.pages.sales_page import SalesPage
    inv, sal, rep = (InventoryPage(page, app.inventory), SalesPage(page, app.sales_ctl),
                     ReportPage(page, app.report))
    for p in (inv, sal, rep):
        p.build()
    return inv, sal, rep


def test_app_builds_with_three_tabs(ui):
    page, tabs, _ = ui
    assert tabs.length == 3 and page.on_resize is not None
    page.width = 1400
    page.on_resize()           # el contenido se limita a 900 px en pantallas grandes
    assert page.added[0].content.controls[0].width == c.MAX_CONTENT_WIDTH


def test_inventory_flow_with_low_stock_alert(app):
    page = FakePage()
    inv, _, _ = _pages(app, page)
    inv.name.value, inv.cost.value, inv.price.value = "Pepsi", "10", "25"
    inv.qty.value, inv.min_stock.value = "1", "5"
    inv.on_add(None)
    assert "Alerta" in snack_text(page)                      # cantidad < mínimo
    assert inv.alert_box.visible and "Pepsi" in inv.alert_text.value
    assert len(inv.rows.controls) == 1
    assert "$10.00" in inv.summary.value                     # valor = costo x cantidad
    inv.select(1)
    assert inv.name.value == "Pepsi" and not inv.update_btn.disabled
    inv.qty.value = "9"
    inv.on_update(None)
    assert not inv.alert_box.visible
    inv.select(1)
    inv.on_delete_click(None)
    assert isinstance(page.dialogs[-1], ft.AlertDialog)
    inv._confirm_delete(page.dialogs[-1])
    assert inv.rows.controls == []


def test_inventory_validation_error_shows_red_snack(app):
    page = FakePage()
    inv, _, _ = _pages(app, page)
    inv.on_add(None)
    assert page.dialogs[-1].bgcolor == "#B91C1C" and "nombre" in snack_text(page)


def test_sale_flow_discount_iva_switch_and_report(app):
    page = FakePage()
    inv, sal, rep = _pages(app, page)
    p = app.inventory.add("Churritos", "6", "12", "10", "3")
    sal.refresh()
    assert sal.empty_text.visible and not sal.order_box.visible
    sal.product.value, sal.qty.value = str(p.id), "8"
    sal.on_add(None)
    assert sal.order_box.visible and not sal.empty_text.visible
    assert sal.l_total.value == "$111.36"                       # 96.00 + 16 % IVA
    sal.iva_switch.value = False                                # Cobrar IVA (16%) apagado
    sal._update_totals(update=False)
    assert sal.l_total.value == "$96.00" and sal.l_iva.value == "$0.00"
    sal.iva_switch.value = True
    # Descuento Global (-%) abre un diálogo
    sal.on_discount(None)
    dlg = page.dialogs[-1]
    dlg.content.value = "abc"
    dlg.actions[2].on_click(None)                               # Aplicar con dato inválido
    assert page.dialogs[-1].bgcolor == "#B91C1C" or isinstance(page.dialogs[-1], ft.SnackBar)
    page.dialogs.remove(dlg)
    sal.on_discount(None)
    dlg = page.dialogs[-1]
    dlg.content.value = "10"
    dlg.actions[2].on_click(None)
    assert sal.discount_btn.content == "Descuento Global (-10%)"
    assert sal.l_discount.value == "-$9.60"
    sal.on_discount(None)
    page.dialogs[-1].actions[1].on_click(None)                  # Quitar
    assert sal.discount_btn.content == "Descuento Global (-%)"
    sal.on_checkout(None)
    assert "Stock bajo" in snack_text(page) and sal.rows.controls == []
    rep.refresh()
    assert rep.heading.value.startswith("RESUMEN DE VENTAS DEL DÍA (")
    texts = [r.controls[1].value for r in rep.metrics.controls]
    assert texts == ["$111.36", "$48.00", "1"]
    assert rep.net.value == "$48.00" and rep.report.units_sold == 8
    assert rep.next_btn.disabled and rep.today_btn.disabled     # ya estamos en hoy
    rep.on_previous(None)
    assert not rep.next_btn.disabled and not rep.today_btn.disabled
    assert rep.report.transactions == 0
    rep.on_next(None)
    assert rep.next_btn.disabled


def test_export_excel_uses_file_picker(app, monkeypatch):
    import io
    from openpyxl import load_workbook
    page = FakePage()
    _, _, rep = _pages(app, page)
    saved = {}

    async def fake_save(page_, filename, data, dialog_title=""):
        saved["name"], saved["data"] = filename, data
        return "C:/tmp/" + filename

    monkeypatch.setattr(c, "save_bytes_file", fake_save)
    rep.refresh()
    asyncio.run(rep.on_export(None))
    assert saved["name"].startswith("reporte_ventas_") and saved["name"].endswith(".xlsx")
    ws = load_workbook(io.BytesIO(saved["data"])).active
    assert ws["A1"].value.startswith("RESUMEN DE VENTAS DEL DÍA")
    assert "guardado" in snack_text(page)


def test_print_and_close_day_buttons_restored(app):
    page = FakePage()
    _, _, rep = _pages(app, page)
    body = rep.build()
    labels = [b.content for row in body.controls if isinstance(row, ft.Row)
              for b in row.controls if hasattr(b, "content") and isinstance(b.content, str)]
    assert "Imprimir informe de ventas" in labels
    assert "Cerrar día y finalizar informe" in labels
    assert "Exportar a Excel" in labels
    assert not any("CSV" in x for x in labels)           # CSV reemplazado por Excel
    rep.refresh()
    assert rep.status.value == ""
    rep.on_close_day(None)
    assert rep.status.value.startswith("DÍA CERRADO")
    rep.on_print(None)
    dlg = page.dialogs[-1]
    assert isinstance(dlg, ft.AlertDialog) and "RESUMEN DE VENTAS" in dlg.content.value
    asyncio.run(rep._copy_and_close(dlg, "texto"))
    assert page.clipboard.value == "texto"

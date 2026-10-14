from decimal import Decimal
from datetime import date

import pytest

from controller.cart import Cart, calculate_totals
from controller.money import format_money, parse_int, parse_money, parse_percent
from model.errors import InsufficientStockError, ValidationError


# ---------- dinero ----------
@pytest.mark.parametrize("text,cents", [("12", 1200), ("12.5", 1250), ("12,50", 1250),
                                        ("$1,200.00", 120000), ("0", 0)])
def test_parse_money(text, cents):
    assert parse_money(text, "el costo") == cents


@pytest.mark.parametrize("bad", ["", "abc", "-5", "nan"])
def test_parse_money_invalid(bad):
    with pytest.raises(ValidationError):
        parse_money(bad, "el costo")


def test_parse_int_and_percent():
    assert parse_int("7", "la cantidad") == 7
    with pytest.raises(ValidationError):
        parse_int("1.5", "la cantidad")
    assert parse_percent("") == 0
    with pytest.raises(ValidationError):
        parse_percent("101")


def test_format_money():
    assert format_money(123456) == "$1,234.56"
    assert format_money(0) == "$0.00"


# ---------- inventario ----------
def test_inventory_add_validates(app):
    with pytest.raises(ValidationError):
        app.inventory.add("", "1", "2", "1", "0")
    with pytest.raises(ValidationError):
        app.inventory.add("A", "x", "2", "1", "0")
    p = app.inventory.add("  Churritos  ", "6", "12", "23", "5")
    assert (p.id, p.name, p.cost_cents, p.price_cents) == (1, "Churritos", 600, 1200)


def test_inventory_update(app):
    p = app.inventory.add("A", "1", "2", "10", "2")
    q = app.inventory.update(p.id, "A", "1", "2", "1", "2")
    assert q.quantity == 1 and q.is_low


# ---------- totales ----------
def test_totals_with_iva_and_discount():
    from model.entities import CartLine
    lines = [CartLine(1, "A", 2, 10000, 6000)]       # subtotal 200.00, costo 120.00
    t = calculate_totals(lines, Decimal("10"))        # desc 20.00 -> base 180.00 -> IVA 28.80
    assert (t.subtotal_cents, t.discount_cents, t.iva_cents, t.total_cents) == \
        (20000, 2000, 2880, 20880)
    assert t.profit_cents == 20000 - 2000 - 12000


# ---------- carrito / venta ----------
def test_cart_merges_lines_and_blocks_overselling(app):
    p = app.inventory.add("A", "5", "10", "3", "1")
    cart = Cart()
    cart.add(p, 2)
    cart.add(p, 1)
    assert len(cart.lines) == 1 and cart.lines[0].quantity == 3
    with pytest.raises(InsufficientStockError):
        cart.add(p, 1)


def test_checkout_updates_stock_and_alerts_low_stock(app):
    p = app.inventory.add("A", "5", "10", "5", "3")
    cart = Cart()
    app.sales_ctl.add_to_cart(cart, str(p.id), "3")
    result = app.sales_ctl.checkout(cart, "0")
    assert app.products.get(p.id).quantity == 2
    assert [x.id for x in result.low_stock] == [p.id]
    assert cart.is_empty
    assert result.totals.total_cents == 3000 + 480      # 30.00 + 16% IVA


def test_checkout_empty_cart_fails(app):
    with pytest.raises(ValidationError):
        app.sales_ctl.checkout(Cart())


def test_checkout_is_atomic_when_stock_changed_meanwhile(app):
    p = app.inventory.add("A", "5", "10", "3", "0")
    cart = Cart()
    app.sales_ctl.add_to_cart(cart, str(p.id), "3")
    app.inventory.update(p.id, "A", "5", "10", "1", "0")   # alguien bajó el stock
    with pytest.raises(InsufficientStockError):
        app.sales_ctl.checkout(cart)
    assert app.products.get(p.id).quantity == 1             # sin cambios
    assert app.report.report(app.report.today()).transactions == 0  # rollback de la venta


# ---------- reporte ----------
def _sell(app, pid, qty, disc="0"):
    cart = Cart()
    app.sales_ctl.add_to_cart(cart, str(pid), str(qty))
    return app.sales_ctl.checkout(cart, disc)


def test_daily_report_totals(app):
    a = app.inventory.add("A", "5", "10", "20", "0")
    b = app.inventory.add("B", "2", "4", "20", "0")
    _sell(app, a.id, 4, "10")     # sub 40.00, desc 4.00, iva 5.76, total 41.76, costo 20.00
    _sell(app, b.id, 5)           # sub 20.00, iva 3.20, total 23.20, costo 10.00
    r = app.report.report(app.report.today())
    assert r.transactions == 2
    assert r.subtotal_cents == 6000 and r.discount_cents == 400 and r.iva_cents == 896
    assert r.total_cents == 6000 - 400 + 896
    assert r.gross_profit_cents == 6000 - 3000
    assert r.net_profit_cents == 3000 - 400
    assert [p.name for p in r.products] == ["A", "B"]


def test_report_survives_deleting_product_and_filters_by_day(app, clock):
    from datetime import datetime
    a = app.inventory.add("A", "5", "10", "20", "0")
    _sell(app, a.id, 1)
    app.inventory.delete(a.id)
    assert app.report.report(app.report.today()).products[0].name == "A"
    assert app.report.report(date(2026, 10, 12)).transactions == 0
    assert app.report.previous_day(date(2026, 10, 13)) == date(2026, 10, 12)


def test_close_day_and_reopen_on_new_sale(app):
    a = app.inventory.add("A", "5", "10", "20", "0")
    day = app.report.today()
    assert app.report.close_day(day).closed is True
    _sell(app, a.id, 1)
    assert app.report.report(day).closed is False


def test_csv_and_text_export(app):
    a = app.inventory.add("A, grande", "5", "10", "20", "0")
    _sell(app, a.id, 2)
    r = app.report.report(app.report.today())
    csv_text = app.report.to_csv(r)
    assert '"A, grande",2,20.00,10.00' in csv_text
    assert "TOTAL NETO GANADO" in app.report.to_text(r)


def test_iva_can_be_turned_off():
    from model.entities import CartLine
    t = calculate_totals([CartLine(1, "A", 1, 10000, 6000)], Decimal("0"), charge_iva=False)
    assert t.iva_cents == 0 and t.total_cents == 10000


def test_checkout_without_iva_is_stored_without_iva(app):
    p = app.inventory.add("A", "5", "10", "5", "0")
    cart = Cart()
    app.sales_ctl.add_to_cart(cart, str(p.id), "2")
    app.sales_ctl.checkout(cart, "", charge_iva=False)
    r = app.report.report(app.report.today())
    assert r.iva_cents == 0 and r.total_cents == 2000 and r.units_sold == 2


def test_xlsx_export_contents(app):
    import io
    from openpyxl import load_workbook
    a = app.inventory.add("Sabritas 300 g", "22", "32", "5", "0")
    _sell(app, a.id, 1)
    data = app.report.to_xlsx(app.report.report(app.report.today()))
    ws = load_workbook(io.BytesIO(data)).active
    cells = {str(c.value): c for row in ws.iter_rows() for c in row if c.value is not None}
    assert "Sabritas 300 g" in cells
    # fila del producto: cantidad, importe y ganancia como NÚMEROS (no texto)
    r = cells["Sabritas 300 g"].row
    assert [ws.cell(r, col).value for col in (2, 3, 4)] == [1, 32.0, 10.0]
    total_row = cells["TOTAL NETO GANADO"].row
    assert ws.cell(total_row, 2).value == 10.0
    ing_row = cells["INGRESOS TOTALES (ventas + IVA − desc.)"].row
    assert ws.cell(ing_row, 2).value == 37.12


def test_next_day_navigation(app):
    from datetime import date
    assert app.report.next_day(date(2026, 10, 12)) == date(2026, 10, 13)
    assert app.report.can_go_next(date(2026, 10, 12)) is True
    assert app.report.can_go_next(date(2026, 10, 13)) is False
